import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "print-ready"
DATA = json.loads((ROOT / "quiz_content.json").read_text(encoding="utf-8"))
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def check_table_geometry(path):
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = etree.fromstring(archive.read(name))
            for table in root.xpath(".//w:tbl", namespaces=NS):
                tbl_w = table.xpath("./w:tblPr/w:tblW", namespaces=NS)
                grid = table.xpath("./w:tblGrid/w:gridCol", namespaces=NS)
                if not tbl_w or not grid:
                    raise AssertionError(f"Missing fixed table geometry in {path.name}:{name}")
                if tbl_w[0].get(f"{{{NS['w']}}}type") != "dxa":
                    raise AssertionError(f"Non-DXA table width in {path.name}:{name}")
                grid_widths = [int(col.get(f"{{{NS['w']}}}w")) for col in grid]
                if sum(grid_widths) != int(tbl_w[0].get(f"{{{NS['w']}}}w")):
                    raise AssertionError(f"Table/grid width mismatch in {path.name}:{name}")
                for row in table.xpath("./w:tr", namespaces=NS):
                    cells = row.xpath("./w:tc", namespaces=NS)
                    for idx, cell in enumerate(cells):
                        tc_w = cell.xpath("./w:tcPr/w:tcW", namespaces=NS)
                        if not tc_w or tc_w[0].get(f"{{{NS['w']}}}type") != "dxa":
                            raise AssertionError(f"Missing DXA cell width in {path.name}:{name}")
                        expected = grid_widths[min(idx, len(grid_widths) - 1)]
                        if int(tc_w[0].get(f"{{{NS['w']}}}w")) != expected:
                            raise AssertionError(f"Cell/grid width mismatch in {path.name}:{name}")


def main():
    report = []
    for quiz in DATA:
        number = quiz["quiz"]
        student_path = OUT / f"Knowledge-Check-{number}-Student.docx"
        key_path = OUT / f"Knowledge-Check-{number}-Key.docx"
        for path in (student_path, key_path):
            assert path.exists() and path.stat().st_size > 20_000, f"Missing or undersized {path}"
            check_table_geometry(path)

        student = Document(student_path)
        student_text = all_text(student)
        assert "INSTRUCTOR COPY" not in student_text
        assert "Answer Key" not in student_text
        prompts = [p.text for p in student.paragraphs if re.match(r"^\d+\. ", p.text)]
        assert len(prompts) == len(quiz["questions"]), (number, len(prompts))
        for question in quiz["questions"]:
            assert question["prompt"] in student_text, (number, question["number"], "prompt")
            for option in question["options"]:
                assert option["text"] in student_text, (number, question["number"], option["text"])
            if question.get("code"):
                assert question["code"] in student_text, (number, question["number"], "code")

        key = Document(key_path)
        key_text = all_text(key)
        assert "INSTRUCTOR COPY" in key_text and "Answer Key" in key_text
        answer_table = key.tables[0]
        assert len(answer_table.rows) == len(quiz["questions"]) + 1
        for question, row in zip(quiz["questions"], answer_table.rows[1:]):
            correct_index, correct = next(
                (idx, option) for idx, option in enumerate(question["options"]) if option["correct"]
            )
            values = [cell.text for cell in row.cells]
            assert values == [str(question["number"]), chr(65 + correct_index), correct["text"]], (
                number,
                question["number"],
                values,
            )

        report.append(
            {
                "quiz": number,
                "questions": len(quiz["questions"]),
                "student_bytes": student_path.stat().st_size,
                "key_bytes": key_path.stat().st_size,
                "geometry": "passed",
                "content": "passed",
            }
        )
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
