import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "print-ready"
DATA = json.loads((ROOT / "quiz_content.json").read_text(encoding="utf-8"))

INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x59, 0x59, 0x59)
LIGHT = "EDEDED"
LIGHTER = "F7F7F7"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, style="nil", color="D9D9D9", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), style)
        if style != "nil":
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fallback, fld_end])
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc, quiz_number, key=False):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5), 16, 8),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5), 12, 6),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78), 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(footer_table, [6840, 2520], indent_dxa=120)
    set_table_borders(footer_table, style="nil")
    left = footer_table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    set_run_font(left.add_run(f"MKTG 411  |  Knowledge Check {quiz_number}" + (" - Key" if key else "")), size=9, color=MUTED)
    right = footer_table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    set_run_font(right.add_run("Page "), size=9, color=MUTED)
    add_page_field(right)

    doc.core_properties.title = f"MKTG 411 Knowledge Check {quiz_number}" + (" - Instructor Key" if key else "")
    doc.core_properties.subject = "Printable in-class knowledge check"
    doc.core_properties.author = "MKTG 411"


def add_student_header(doc, quiz_number, total_points):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(1)
    set_run_font(kicker.add_run("MKTG 411  |  FALL 2026"), size=9, bold=True, color=MUTED)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run(f"Knowledge Check {quiz_number}"), size=22, bold=True, color=INK)

    info = doc.add_table(rows=2, cols=2)
    set_table_geometry(info, [5400, 3960], indent_dxa=120)
    set_table_borders(info, style="single", color="D9D9D9", size="4")
    fields = (
        (0, 0, "Name:  __________________________________________"),
        (0, 1, "Section:  __________"),
        (1, 0, "Date:  ____________________"),
        (1, 1, f"Score:  ______ / {total_points}"),
    )
    for row, col, text in fields:
        p = info.cell(row, col).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=10.5, bold=(col == 1 and row == 1))

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)

    instructions = doc.add_paragraph()
    instructions.paragraph_format.space_after = Pt(10)
    instructions.paragraph_format.keep_with_next = True
    set_run_font(instructions.add_run("Instructions: "), size=10.5, bold=True)
    set_run_font(
        instructions.add_run(
            "Complete this written quiz individually during class. Circle the best answer for each question. "
            "You may use printed or handwritten course materials. No devices or technological assistance are permitted."
        ),
        size=10.5,
    )


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table, style="single", color="D9D9D9", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHTER)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(code), name="Consolas", size=8.5, color=INK)
    return table


def add_question(doc, question):
    prompt = doc.add_paragraph()
    prompt.paragraph_format.space_before = Pt(5)
    prompt.paragraph_format.space_after = Pt(3)
    prompt.paragraph_format.keep_with_next = True
    set_run_font(prompt.add_run(f"{question['number']}. "), size=10.5, bold=True)
    set_run_font(prompt.add_run(question["prompt"]), size=10.5)

    if question.get("code"):
        add_code_block(doc, question["code"])

    options = doc.add_table(rows=len(question["options"]), cols=2)
    set_table_geometry(options, [540, 8820], indent_dxa=120)
    set_table_borders(options, style="nil")
    for idx, option in enumerate(question["options"]):
        row = options.rows[idx]
        prevent_row_split(row)
        label = row.cells[0].paragraphs[0]
        label.paragraph_format.space_after = Pt(0)
        label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(label.add_run(f"{chr(65 + idx)}."), size=10.5, bold=True)
        text = row.cells[1].paragraphs[0]
        text.paragraph_format.space_after = Pt(0)
        set_run_font(text.add_run(option["text"]), size=10.5)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.line_spacing = 0.5


def build_student(quiz):
    doc = Document()
    configure_document(doc, quiz["quiz"], key=False)
    add_student_header(doc, quiz["quiz"], len(quiz["questions"]))
    for question in quiz["questions"]:
        add_question(doc, question)
    path = OUTPUT / f"Knowledge-Check-{quiz['quiz']}-Student.docx"
    doc.save(path)
    return path


def build_key(quiz):
    doc = Document()
    configure_document(doc, quiz["quiz"], key=True)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(1)
    set_run_font(kicker.add_run("MKTG 411  |  INSTRUCTOR COPY"), size=9, bold=True, color=MUTED)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run(f"Knowledge Check {quiz['quiz']} - Answer Key"), size=22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    set_run_font(subtitle.add_run(f"{len(quiz['questions'])} points total | One point per question"), size=10.5, italic=True, color=MUTED)

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [720, 1080, 7560], indent_dxa=120)
    set_table_borders(table, style="single", color="BFBFBF", size="4")
    headers = ("#", "Answer", "Correct response")
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=10, bold=True)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for question in quiz["questions"]:
        correct_idx, correct = next((idx, opt) for idx, opt in enumerate(question["options"]) if opt["correct"])
        cells = table.add_row().cells
        values = (str(question["number"]), chr(65 + correct_idx), correct["text"])
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=10, bold=(idx == 1))
        prevent_row_split(table.rows[-1])
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=70, bottom=70, start=120, end=120)

    path = OUTPUT / f"Knowledge-Check-{quiz['quiz']}-Key.docx"
    doc.save(path)
    return path


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    created = []
    for quiz in DATA:
        created.extend([build_student(quiz), build_key(quiz)])
    for path in created:
        print(path)


if __name__ == "__main__":
    main()
